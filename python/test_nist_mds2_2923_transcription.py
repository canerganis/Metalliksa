"""Verify the pinned NIST mds2-2923 IN718 summary transcription."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1] / "data/benchmark/nist-mds2-2923-in718/official"


def main() -> None:
    manifest = json.loads((ROOT / "manifest.json").read_text(encoding="utf-8"))
    for artifact in manifest["files"]:
        content = (ROOT / artifact["path"]).read_bytes()
        assert len(content) == artifact["bytes"], artifact["path"]
        assert hashlib.sha256(content).hexdigest() == artifact["sha256"], artifact["path"]

    workbook = load_workbook(ROOT / "Master_TrackList_Measurements.xlsx", read_only=True, data_only=True)
    summary = list(workbook["Summary"].iter_rows(values_only=True))
    headers = {name: index for index, name in enumerate(summary[0]) if name is not None}
    expected = {
        "machine": "Machine",
        "beamDiameter_um": "Estimated D4σ Spot Diameter (µm)",
        "laserPower_W": "Laser Power (W)",
        "scanSpeed_mm_s": "Scan Speed (mm/s)",
        "observationCount": "no. of measurements",
        "meanWidth_um": "Average Width (µm)",
        "widthStdDev_um": "Std. dev. Width (µm)",
        "expandedWidthUncertainty_k2_um": "Width: Combined, Expanded uncertainty (µm); U (k=2)",
        "meanDepth_um": "Average Depth (µm)",
        "depthStdDev_um": "Std. dev. Depth (µm)",
        "expandedDepthUncertainty_k2_um": "Depth: Combined, Expanded uncertainty (µm); U (k-2)",
    }
    rows = [row for row in summary[1:] if row[headers["Material"]] == "IN718"]
    assert len(rows) == len(manifest["measurements"]) == 6
    for source, recorded in zip(rows, manifest["measurements"], strict=True):
        for key, header in expected.items():
            value = source[headers[header]]
            if isinstance(value, (float, int)) and not isinstance(value, bool):
                assert abs(float(value) - float(recorded[key])) <= 1e-9, (key, value, recorded[key])
            else:
                assert value == recorded[key], (key, value, recorded[key])
        assert recorded["beamDiameterSource"] == ("estimated" if recorded["machine"] == "EOS M290" else "measured")
        assert recorded["meanDepth_um"] / (recorded["beamDiameter_um"] / 2) > 3

    data = list(workbook["Data"].iter_rows(values_only=True))
    assert sum(row[5] == "IN718" for row in data[1:]) == 72
    print("NIST mds2-2923 IN718 transcription: 6 groups, 72 rows, hashes and U(k=2) values PASS")


if __name__ == "__main__":
    main()
